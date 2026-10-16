from dataclasses import dataclass
import hashlib
from pathlib import Path
import re
from uuid import NAMESPACE_URL, uuid5
import zipfile


@dataclass
class ParsedSection:
    text: str
    page: int
    section: str = ""


def validate_file(path: Path, extension: str):
    if extension == ".pdf":
        with path.open("rb") as file:
            if file.read(5) != b"%PDF-":
                raise ValueError("The uploaded file is not a PDF")
    elif extension == ".docx":
        if not zipfile.is_zipfile(path):
            raise ValueError("The uploaded file is not a DOCX")
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(x.file_size for x in entries) > 100 * 1024 * 1024:
                raise ValueError("DOCX archive expands beyond the processing limit")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("DOCX is missing its document XML")
            if any(x.flag_bits & 1 or x.file_size > max(1, x.compress_size) * 300 for x in entries):
                raise ValueError("Encrypted or excessively compressed DOCX is not accepted")
    elif extension not in {".txt", ".md"}:
        raise ValueError("Allowed file types: PDF, TXT, Markdown and DOCX")


def clean(text: str):
    text = text.replace("\x00", "").replace("\r\n", "\n")
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", text)).strip()


def parse_file(path: Path, extension: str) -> list[ParsedSection]:
    validate_file(path, extension)
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported")
        if len(reader.pages) > 500:
            raise ValueError("PDF exceeds the 500-page processing limit")
        sections = [ParsedSection(clean(page.extract_text() or ""), i + 1) for i, page in enumerate(reader.pages)]
    elif extension == ".docx":
        from docx import Document
        document = Document(path)
        sections, heading, text = [], "", []
        for paragraph in document.paragraphs:
            if paragraph.style.name.startswith("Heading") and text:
                sections.append(ParsedSection(clean("\n".join(text)), 0, heading))
                text = []
            if paragraph.style.name.startswith("Heading"):
                heading = paragraph.text
            text.append(paragraph.text)
        for table in document.tables:
            text.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        if text:
            sections.append(ParsedSection(clean("\n".join(text)), 0, heading))
    else:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
        if extension == ".md":
            sections, heading, content = [], "", []
            for line in text.splitlines():
                if line.startswith("#") and content:
                    sections.append(ParsedSection(clean("\n".join(content)), 0, heading))
                    content = []
                if line.startswith("#"):
                    heading = line.lstrip("# ")[:255]
                content.append(line)
            if content:
                sections.append(ParsedSection(clean("\n".join(content)), 0, heading))
        else:
            sections = [ParsedSection(clean(text), 0)]
    sections = [s for s in sections if s.text]
    if not sections:
        raise ValueError("No extractable text found; scanned PDFs require an OCR preprocessing step")
    if sum(len(s.text) for s in sections) > 3_000_000:
        raise ValueError("Extracted text exceeds the processing limit")
    return sections


def chunk_sections(document_id: str, sections: list[ParsedSection], size=1500, overlap=200):
    if size <= overlap or overlap < 0:
        raise ValueError("Chunk size must exceed nonnegative overlap")
    chunks = []
    for section in sections:
        start = 0
        while start < len(section.text):
            end = min(start + size, len(section.text))
            if end < len(section.text):
                boundary = max(section.text.rfind("\n", start + size // 2, end), section.text.rfind(". ", start + size // 2, end))
                if boundary > start:
                    end = boundary + 1
            content = section.text[start:end].strip()
            if content:
                position = len(chunks)
                digest = hashlib.sha256(content.encode()).hexdigest()
                chunks.append({"id": str(uuid5(NAMESPACE_URL, f"{document_id}:{position}:{digest}")), "position": position,
                               "content": content, "page": section.page, "section": section.section[:255], "metadata_json": {"sha256": digest}})
            if end == len(section.text):
                break
            start = max(start + 1, end - overlap)
    return chunks
